import os
import re
import tkinter as tk
import threading
import shutil

from tkinter import ttk
from tkinter.filedialog import askopenfilename, askdirectory

from openpyxl import load_workbook
from openpyxl import load_workbook

# ---------- Configurações ----------
PREFIXO_ARQUIVO = ""

ABAS_IGNORADAS_POR_NOME_RAW = {"Instruções", ""}
ABAS_IGNORADAS_POR_NOME = {s.strip().lower() for s in ABAS_IGNORADAS_POR_NOME_RAW}

IGNORAR_PRIMEIRAS_N_ABAS = 0


# ---------- Funções auxiliares ----------
def normalizar_nome_arquivo(nome: str) -> str:
    nome = re.sub(r'[\\/:*?"<>|]', "_", (nome or "")).strip()
    return nome[:150]

def normalizar_nome_aba(nome: str) -> str:
    return (nome or "").strip().lower()

def extrair_aba_completa(caminho_original, caminho_saida, nome_aba):
    # copia o arquivo inteiro
    shutil.copy(caminho_original, caminho_saida)

    # abre a cópia
    wb = load_workbook(caminho_saida)

    # remove todas as abas exceto a desejada
    for aba in wb.sheetnames:
        if aba != nome_aba:
            del wb[aba]

    # salva
    wb.save(caminho_saida)
    wb.close()

# ---------------- UI ----------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Spreadsheet Tab Extractor")
        self.geometry("560x300")

        self.caminho_excel = None
        self.pasta_saida = None

        self.create_widgets()

    def create_widgets(self):
        ttk.Label(self, text="Dividir abas em arquivos", font=("Segoe UI", 12, "bold")).pack(pady=10)

        frame = ttk.Frame(self)
        frame.pack()

        ttk.Button(frame, text="📎 Anexar", command=self.on_anexar).grid(row=0, column=0, padx=5)

        self.btn = ttk.Button(
            frame,
            text="▶️ Processar",
            state=tk.DISABLED,
            command=self.on_thread
        )
        self.btn.grid(row=0, column=1, padx=5)

        self.logbox = tk.Text(self, height=12)
        self.logbox.pack(padx=10, pady=10)

    def log(self, msg):
        self.logbox.insert(tk.END, msg + "\n")
        self.logbox.see(tk.END)

    def on_anexar(self):
        path = askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm *.xls")])
        if path:
            self.caminho_excel = path
            self.log(f"Arquivo: {path}")
            self.btn.config(state=tk.NORMAL)

    def on_thread(self):
        pasta = askdirectory(title="Selecione diretório")
        if not pasta:
            return

        self.pasta_saida = pasta

        threading.Thread(target=self.processar).start()

    def processar(self):
        self.log("Lendo abas...")

        wb = load_workbook(self.caminho_excel, data_only=True)

        for idx, nome in enumerate(wb.sheetnames, start=1):

            nome_norm = normalizar_nome_aba(nome)

            if idx <= IGNORAR_PRIMEIRAS_N_ABAS:
                continue

            if nome_norm in ABAS_IGNORADAS_POR_NOME:
                self.log(f"↷ Ignorada: {nome}")
                continue

            try:
                self.log(f"⚙️ Processando: {nome}")
                self.update()

                caminho_saida = os.path.join(
                    self.pasta_saida,
                    f"{PREFIXO_ARQUIVO}{normalizar_nome_arquivo(nome)}.xlsx"
                )

                extrair_aba_completa(
                    self.caminho_excel,
                    caminho_saida,
                    nome
                )

                self.log(f"✅ {nome}")
                self.update()

            except Exception as e:
                self.log(f"❌ Erro em {nome}: {e}")
                continue

        wb.close()
        self.log("✅ Finalizado!")

def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
